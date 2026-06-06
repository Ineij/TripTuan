#!/usr/bin/env python3
"""Sync generated POI image URLs back into a CSV, with optional XLSX output.

Expected CSV columns:
  poi_id,name,type,city,source_image_url

Example:
  python sync_generated_image_links.py \
    --csv ./sz_poi.csv \
    --scene sz \
    --public-base-url http://127.0.0.1:8000 \
    --xlsx
"""
from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path
from urllib.parse import quote


ROOT = Path(os.getenv("TRIPTUAN_ROOT", Path.cwd())).resolve()
OUTPUT_DIR = ROOT / "backend" / "static" / "generated" / "pois" / "csv-preview"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv", required=True, help="Input CSV path")
    parser.add_argument("--scene", default="", help="Scene key, e.g. sz or bj. Defaults from city column")
    parser.add_argument(
        "--public-base-url",
        default=os.getenv("POI_IMAGE_PUBLIC_BASE_URL", "http://127.0.0.1:8000"),
        help="URL prefix written into generated_image_url",
    )
    parser.add_argument(
        "--output-csv",
        default="",
        help="Output CSV path. Defaults to csv-preview/<input>_with_generated_images.csv",
    )
    parser.add_argument("--xlsx", action="store_true", help="Also write an XLSX file if openpyxl is installed")
    args = parser.parse_args()

    input_csv = Path(args.csv)
    output_csv = Path(args.output_csv) if args.output_csv else OUTPUT_DIR / f"{input_csv.stem}_with_generated_images.csv"
    rows, fieldnames = read_csv(input_csv)
    if "generated_image_url" not in fieldnames:
        fieldnames.append("generated_image_url")

    filled = 0
    for row in rows:
        scene = args.scene or infer_scene(row.get("city", ""))
        candidate = OUTPUT_DIR / (
            f"{scene}_{safe_slug(row.get('poi_id', ''))}_{safe_slug(row.get('name', ''))}_ins.png"
        )
        if candidate.exists():
            row["generated_image_url"] = public_url_for(candidate, args.public_base_url)
            filled += 1
        else:
            row.setdefault("generated_image_url", "")

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"rows={len(rows)} generated_links={filled}", flush=True)
    print(f"CSV -> {output_csv}", flush=True)

    if args.xlsx:
        output_xlsx = output_csv.with_suffix(".xlsx")
        write_xlsx(rows, fieldnames, output_xlsx)
        print(f"XLSX -> {output_xlsx}", flush=True)
    return 0


def read_csv(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        return list(reader), list(reader.fieldnames or [])


def infer_scene(city: str) -> str:
    text = city.lower()
    if "北京" in city or "beijing" in text:
        return "bj"
    if "深圳" in city or "shenzhen" in text:
        return "sz"
    return "poi"


def safe_slug(text: str, fallback: str = "poi") -> str:
    slug = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", (text or "")).strip("_")
    return slug[:40] or fallback


def public_url_for(path: Path, public_base_url: str) -> str:
    rel = "/" + quote(path.relative_to(ROOT / "backend").as_posix(), safe="/")
    return f"{public_base_url.rstrip('/')}{rel}"


def write_xlsx(rows: list[dict[str, str]], fieldnames: list[str], output_xlsx: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("openpyxl is required for --xlsx. Install it with: python -m pip install openpyxl") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "POI images"
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(field, "") for field in fieldnames])

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if "generated_image_url" in fieldnames:
        link_col = fieldnames.index("generated_image_url") + 1
        for row_idx in range(2, len(rows) + 2):
            cell = sheet.cell(row_idx, link_col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    widths = {
        "poi_id": 10,
        "name": 34,
        "type": 12,
        "city": 12,
        "source_image_url": 52,
        "generated_image_url": 84,
    }
    for idx, field in enumerate(fieldnames, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = widths.get(field, 18)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    workbook.save(output_xlsx)


if __name__ == "__main__":
    raise SystemExit(main())
